from pathlib import Path
from datetime import datetime, timezone

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    HRFlowable,
    KeepTogether,
    ListFlowable,
    ListItem,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
)

ROOT = Path(__file__).resolve().parents[1]
DOCX_OUT = ROOT / "output" / "CV-Juan-Cabello.docx"
PDF_OUT = ROOT / "public" / "cv-juan-cabello.pdf"

NAVY = "102033"
BLUE = "0067C5"
MUTED = "4B5F73"
LIGHT = "E6F2FF"
WHITE = "FFFFFF"


def shade_paragraph(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def add_hyperlink(paragraph, text: str, url: str, color: str = BLUE) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color_element = OxmlElement("w:color")
    color_element.set(qn("w:val"), color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([color_element, underline])
    run.append(properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def docx_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(BLUE)
    run.font.letter_spacing = Pt(0.7)


def add_docx_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.first_line_indent = Inches(-0.12)
    run = paragraph.add_run(text)
    run.font.size = Pt(9.3)
    run.font.color.rgb = RGBColor.from_string(NAVY)


def build_docx() -> None:
    DOCX_OUT.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.page_height = Inches(11.69)
    section.page_width = Inches(8.27)
    section.top_margin = Inches(0.52)
    section.bottom_margin = Inches(0.52)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.3)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.04

    document.core_properties.title = "CV — Juan Xavier Cabello Salirrosas"
    document.core_properties.subject = "Currículum profesional"
    document.core_properties.author = "Juan Xavier Cabello Salirrosas"
    document.core_properties.keywords = "matemática, docencia, desarrollo web, educación STEM"
    document.core_properties.comments = "Generado para uso profesional."
    document.core_properties.created = datetime(2026, 10, 6, tzinfo=timezone.utc)
    document.core_properties.modified = datetime(2026, 10, 6, tzinfo=timezone.utc)

    banner = document.add_paragraph()
    banner.alignment = WD_ALIGN_PARAGRAPH.LEFT
    banner.paragraph_format.space_after = Pt(0)
    banner.paragraph_format.left_indent = Inches(0.16)
    banner.paragraph_format.right_indent = Inches(0.16)
    banner.paragraph_format.space_before = Pt(0)
    shade_paragraph(banner, NAVY)
    run = banner.add_run("JUAN XAVIER CABELLO SALIRROSAS")
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor.from_string(WHITE)

    role = document.add_paragraph()
    role.paragraph_format.left_indent = Inches(0.16)
    role.paragraph_format.right_indent = Inches(0.16)
    role.paragraph_format.space_after = Pt(0)
    shade_paragraph(role, NAVY)
    r = role.add_run("Educador Matemático · Desarrollador Web · Creador de Productos Educativos")
    r.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(140, 205, 255)

    contact = document.add_paragraph()
    contact.paragraph_format.left_indent = Inches(0.16)
    contact.paragraph_format.right_indent = Inches(0.16)
    contact.paragraph_format.space_after = Pt(8)
    shade_paragraph(contact, NAVY)
    cr = contact.add_run("Lima, Perú · Remoto   |   +51 925 475 034   |   ")
    cr.font.size = Pt(8.7)
    cr.font.color.rgb = RGBColor.from_string(WHITE)
    add_hyperlink(contact, "juan@cabellosalirrosas.com", "mailto:juan@cabellosalirrosas.com", "8BCDFF")

    links = document.add_paragraph()
    links.paragraph_format.space_after = Pt(7)
    add_hyperlink(links, "Portafolio", "https://juan.cabellosalirrosas.com/")
    links.add_run("  ·  ")
    add_hyperlink(links, "GitHub", "https://github.com/juanxaviercasa")
    links.add_run("  ·  ")
    add_hyperlink(links, "LinkedIn", "https://www.linkedin.com/in/xaviercabello/")
    links.add_run("  ·  ")
    add_hyperlink(links, "Perfil docente", "https://www.tusclases.pe/profesores/juan-xavier-cabello-salirrosas.htm")

    docx_heading(document, "Perfil profesional")
    profile = document.add_paragraph(
        "Bachiller en Matemática y docente con experiencia en educación secundaria y preparación preuniversitaria. "
        "Complemento la práctica pedagógica con desarrollo web para crear laboratorios visuales, plataformas de evaluación "
        "y herramientas digitales orientadas a resolver necesidades reales. Trabajo con especial atención a claridad, "
        "accesibilidad, diseño responsivo y aprendizaje activo."
    )
    profile.paragraph_format.space_after = Pt(5)

    docx_heading(document, "Competencias clave")
    skill_items = [
        ("Educación", "Didáctica matemática · planificación · evaluación"),
        ("Frontend", "React · TypeScript · JavaScript · HTML · CSS"),
        ("Visualización", "Canvas · WebGL · Three.js · KaTeX"),
        ("Producto", "UX/UI · responsive · accesibilidad · SEO"),
        ("Backend", "Python · FastAPI · PostgreSQL · Docker"),
        ("Colaboración", "Git · documentación · trabajo remoto"),
    ]
    for title, detail in skill_items:
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Inches(0.12)
        p.paragraph_format.right_indent = Inches(0.12)
        shade_paragraph(p, "EDF3F9")
        title_run = p.add_run(title + "\n")
        title_run.bold = True
        title_run.font.color.rgb = RGBColor.from_string(BLUE)
        detail_run = p.add_run(detail)
        detail_run.font.size = Pt(8.2)
        detail_run.font.color.rgb = RGBColor.from_string(MUTED)

    docx_heading(document, "Experiencia profesional")
    roles = [
        ("2026", "Docente de Aritmética y Razonamiento Matemático", "CEPREMUNI Ate", "Ciclos de verano y semestrales; niveles intermedio y avanzado en Centro Cultural, Ollantaytambo y Huaycán."),
        ("2025", "Docente de Álgebra", "CEPREMUNI Ate", "Preparación semestral en sedes de Ate, Huaycán y Ollantaytambo."),
        ("2025", "Docente de Física, Química y Biología", "I.E.P. Jesús de Nazareth–Belén", "Enseñanza a estudiantes de 3.° a 5.° de secundaria."),
        ("2020–2022", "Docente de Matemáticas y Ciencias", "Academia Pre Policial Fuerza Delta", "Razonamiento Matemático, Aritmética, Geometría, Física, Química y Biología en sedes de Lima Este."),
        ("2018", "Docente de Ciencias", "I.E.P. Jesús de Nazareth–Belén", "Física, Química y Biología para 1.° a 5.° de secundaria."),
    ]
    for period, title, org, detail in roles:
        p = document.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(1)
        pr = p.add_run(f"{period}  |  {title}")
        pr.bold = True
        pr.font.size = Pt(9.5)
        pr.font.color.rgb = RGBColor.from_string(NAVY)
        org_p = document.add_paragraph()
        org_p.paragraph_format.space_after = Pt(1)
        org_r = org_p.add_run(org)
        org_r.bold = True
        org_r.font.size = Pt(8.7)
        org_r.font.color.rgb = RGBColor.from_string(BLUE)
        detail_p = document.add_paragraph(detail)
        detail_p.paragraph_format.space_after = Pt(2)
        detail_p.runs[0].font.size = Pt(8.7)
        detail_p.runs[0].font.color.rgb = RGBColor.from_string(MUTED)

    document.add_page_break()
    docx_heading(document, "Proyectos seleccionados")
    projects = [
        ("Mundos Simulados", "Laboratorio web de física computacional y sistemas dinámicos con Canvas/WebGL.", "https://mundossimulados.online", "https://github.com/juanxaviercasa/mundos-simulados"),
        ("Transformaciones Geométricas", "Visualizador 2D/3D de rotaciones, homotecias y transformaciones matriciales en TypeScript.", "https://transformacionesgeometricas.sistemazenit.com/", "https://github.com/juanxaviercasa/transformaciones_geometricas"),
        ("Tabla Periódica Interactiva", "Explorador responsivo de los 118 elementos, filtros y tendencias periódicas.", "https://tablaperiodica.sistemazenit.com/", "https://github.com/juanxaviercasa/tabla-periodica"),
        ("Nube para Pymes", "Colección de calculadoras, generadores y herramientas web para pequeños negocios.", "https://nubeparapymes.online/", "https://github.com/juanxaviercasa/nube-para-pymes"),
        ("Zenit AI Tutor", "Prototipo educativo full stack con FastAPI, Next.js, PostgreSQL/pgvector y RAG con fuentes.", None, "https://github.com/juanxaviercasa/zenit-ai-tutor"),
    ]
    for title, description, live, repo in projects:
        p = document.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor.from_string(NAVY)
        d = document.add_paragraph(description)
        d.paragraph_format.space_after = Pt(1)
        d.runs[0].font.size = Pt(8.8)
        d.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
        l = document.add_paragraph()
        l.paragraph_format.space_after = Pt(2)
        if live:
            add_hyperlink(l, "Ver proyecto", live)
            l.add_run("  ·  ")
        add_hyperlink(l, "Ver código", repo)

    docx_heading(document, "Formación")
    add_docx_bullet(document, "Bachiller en Matemática — Universidad Nacional de Educación Enrique Guzmán y Valle, La Cantuta.")
    add_docx_bullet(document, "Desarrollador Web Full Stack — Make It Real Camp.")
    add_docx_bullet(document, "Programación desde Cero — Egg Cooperation.")

    docx_heading(document, "Enfoque de trabajo")
    add_docx_bullet(document, "Traduzco contenidos complejos en interfaces y explicaciones que priorizan comprensión y autonomía.")
    add_docx_bullet(document, "Diseño primero para móvil, teclado y lectura clara; incorporo animación solo cuando comunica un cambio.")
    add_docx_bullet(document, "Documento decisiones técnicas, valido enlaces y evito métricas o afirmaciones no verificables.")

    footer = document.sections[-1].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("Juan Xavier Cabello Salirrosas · CV actualizado en octubre de 2026")
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor.from_string(MUTED)

    document.save(DOCX_OUT)


def pdf_link(text: str, url: str) -> str:
    return f'<link href="{url}" color="#{BLUE}"><u>{text}</u></link>'


def build_pdf() -> None:
    PDF_OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        str(PDF_OUT),
        pagesize=A4,
        rightMargin=16 * mm,
        leftMargin=16 * mm,
        topMargin=14 * mm,
        bottomMargin=13 * mm,
        title="CV — Juan Xavier Cabello Salirrosas",
        author="Juan Xavier Cabello Salirrosas",
        subject="Currículum profesional",
    )
    styles = getSampleStyleSheet()
    title = ParagraphStyle("TitleCV", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=22, leading=24, textColor=colors.HexColor("#FFFFFF"), spaceAfter=3)
    role = ParagraphStyle("Role", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=10.5, leading=13, textColor=colors.HexColor("#8BCDFF"), spaceAfter=4)
    contact = ParagraphStyle("Contact", parent=styles["Normal"], fontSize=8.8, leading=11, textColor=colors.HexColor("#FFFFFF"))
    heading = ParagraphStyle("Heading", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=10.5, leading=13, textColor=colors.HexColor(f"#{BLUE}"), spaceBefore=8, spaceAfter=4, uppercase=True)
    body = ParagraphStyle("Body", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.2, leading=12.3, textColor=colors.HexColor(f"#{NAVY}"), spaceAfter=3)
    body_muted = ParagraphStyle("BodyMuted", parent=body, fontSize=8.7, leading=11.5, textColor=colors.HexColor(f"#{MUTED}"))
    item_title = ParagraphStyle("ItemTitle", parent=body, fontName="Helvetica-Bold", fontSize=9.6, leading=12, spaceBefore=3, spaceAfter=0)
    link_style = ParagraphStyle("Links", parent=body, fontSize=8.5, leading=11, textColor=colors.HexColor(f"#{BLUE}"), spaceAfter=2)

    story = []
    header_table_data = [[Paragraph("JUAN XAVIER CABELLO SALIRROSAS", title)], [Paragraph("Educador Matemático · Desarrollador Web · Creador de Productos Educativos", role)], [Paragraph(f'Lima, Perú · Remoto &nbsp;|&nbsp; +51 925 475 034 &nbsp;|&nbsp; {pdf_link("juan@cabellosalirrosas.com", "mailto:juan@cabellosalirrosas.com")}', contact)]]
    from reportlab.platypus import Table, TableStyle
    header = Table(header_table_data, colWidths=[178 * mm], hAlign="LEFT")
    header.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(f"#{NAVY}")), ("LEFTPADDING", (0, 0), (-1, -1), 9), ("RIGHTPADDING", (0, 0), (-1, -1), 9), ("TOPPADDING", (0, 0), (-1, 0), 7), ("BOTTOMPADDING", (0, 0), (-1, 0), 1), ("TOPPADDING", (0, 1), (-1, -1), 1), ("BOTTOMPADDING", (0, -1), (-1, -1), 7)]))
    story.extend([header, Spacer(1, 5), Paragraph(f'{pdf_link("Portafolio", "https://juan.cabellosalirrosas.com/")} · {pdf_link("GitHub", "https://github.com/juanxaviercasa")} · {pdf_link("LinkedIn", "https://www.linkedin.com/in/xaviercabello/")} · {pdf_link("Perfil docente", "https://www.tusclases.pe/profesores/juan-xavier-cabello-salirrosas.htm")}', link_style)])
    story.extend([Paragraph("PERFIL PROFESIONAL", heading), Paragraph("Bachiller en Matemática y docente con experiencia en educación secundaria y preparación preuniversitaria. Complemento la práctica pedagógica con desarrollo web para crear laboratorios visuales, plataformas de evaluación y herramientas digitales orientadas a resolver necesidades reales. Trabajo con especial atención a claridad, accesibilidad, diseño responsivo y aprendizaje activo.", body)])

    story.append(Paragraph("COMPETENCIAS CLAVE", heading))
    skill_data = []
    skills = [("Educación", "Didáctica matemática · planificación · evaluación"), ("Frontend", "React · TypeScript · JavaScript · HTML · CSS"), ("Visualización", "Canvas · WebGL · Three.js · KaTeX"), ("Producto", "UX/UI · responsive · accesibilidad · SEO"), ("Backend", "Python · FastAPI · PostgreSQL · Docker"), ("Colaboración", "Git · documentación · trabajo remoto")]
    for row_start in range(0, len(skills), 3):
        row = []
        for name, detail in skills[row_start:row_start + 3]:
            row.append(Paragraph(f'<b><font color="#{BLUE}">{name}</font></b><br/><font size="8" color="#{MUTED}">{detail}</font>', body))
        skill_data.append(row)
    skill_table = Table(skill_data, colWidths=[59 * mm] * 3, hAlign="LEFT")
    skill_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#EDF3F9")), ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD8E6")), ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD8E6")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
    story.append(skill_table)

    story.append(Paragraph("EXPERIENCIA PROFESIONAL", heading))
    roles = [("2026", "Docente de Aritmética y Razonamiento Matemático", "CEPREMUNI Ate", "Ciclos de verano y semestrales; niveles intermedio y avanzado en Centro Cultural, Ollantaytambo y Huaycán."), ("2025", "Docente de Álgebra", "CEPREMUNI Ate", "Preparación semestral en sedes de Ate, Huaycán y Ollantaytambo."), ("2025", "Docente de Física, Química y Biología", "I.E.P. Jesús de Nazareth–Belén", "Enseñanza a estudiantes de 3.° a 5.° de secundaria."), ("2020–2022", "Docente de Matemáticas y Ciencias", "Academia Pre Policial Fuerza Delta", "Razonamiento Matemático, Aritmética, Geometría, Física, Química y Biología en sedes de Lima Este."), ("2018", "Docente de Ciencias", "I.E.P. Jesús de Nazareth–Belén", "Física, Química y Biología para 1.° a 5.° de secundaria.")]
    for period, job, org, detail in roles:
        story.append(KeepTogether([Paragraph(f"<b>{period}</b> &nbsp;|&nbsp; {job}", item_title), Paragraph(f'<b><font color="#{BLUE}">{org}</font></b> — {detail}', body_muted)]))

    story.append(Paragraph("FORMACIÓN", heading))
    education = ["Bachiller en Matemática — Universidad Nacional de Educación Enrique Guzmán y Valle, La Cantuta.", "Desarrollador Web Full Stack — Make It Real Camp.", "Programación desde Cero — Egg Cooperation."]
    story.append(ListFlowable([ListItem(Paragraph(item, body), leftIndent=10) for item in education], bulletType="bullet", start="circle", leftIndent=14, bulletFontName="Helvetica", bulletFontSize=6))

    story.append(Paragraph("ENFOQUE DE TRABAJO", heading))
    approach = ["Traduzco contenidos complejos en interfaces y explicaciones que priorizan comprensión y autonomía.", "Diseño primero para móvil, teclado y lectura clara; incorporo animación solo cuando comunica un cambio.", "Documento decisiones técnicas, valido enlaces y evito métricas o afirmaciones no verificables."]
    story.append(ListFlowable([ListItem(Paragraph(item, body), leftIndent=10) for item in approach], bulletType="bullet", leftIndent=14, bulletFontSize=6))

    story.append(PageBreak())
    story.append(Paragraph("PROYECTOS SELECCIONADOS", heading))
    projects = [("Mundos Simulados", "Laboratorio de física computacional y sistemas dinámicos. Permite variar gravedad, fricción, masa y carga para observar resultados en Canvas/WebGL.", "https://mundossimulados.online", "https://github.com/juanxaviercasa/mundos-simulados"), ("Transformaciones Geométricas", "Visualizador 2D/3D de rotaciones, homotecias y transformaciones matriciales. Desarrollado en TypeScript con controles interactivos.", "https://transformacionesgeometricas.sistemazenit.com/", "https://github.com/juanxaviercasa/transformaciones_geometricas"), ("Tabla Periódica Interactiva", "Explorador responsivo de los 118 elementos con búsqueda, filtros y mapas visuales de propiedades periódicas.", "https://tablaperiodica.sistemazenit.com/", "https://github.com/juanxaviercasa/tabla-periodica"), ("Nube para Pymes", "Colección de calculadoras, generadores de documentos y herramientas de productividad para pequeños negocios, sin registro obligatorio.", "https://nubeparapymes.online/", "https://github.com/juanxaviercasa/nube-para-pymes"), ("Rumbo San Marcos", "Plataforma de evaluación diagnóstica que relaciona carrera, resultados, brechas y prioridades de estudio para postulantes a la UNMSM.", "https://sanmarcos.sistemazenit.com/", "https://github.com/juanxaviercasa/rumbo-san-marcos"), ("Zenit AI Tutor", "Prototipo educativo full stack con FastAPI, Next.js, PostgreSQL/pgvector y RAG con fuentes, validación e incertidumbre explícita.", None, "https://github.com/juanxaviercasa/zenit-ai-tutor")]
    for name, description, live, repo in projects:
        links = (pdf_link("Ver proyecto", live) + " · " if live else "") + pdf_link("Ver código", repo)
        story.append(KeepTogether([Paragraph(name, item_title), Paragraph(description, body_muted), Paragraph(links, link_style)]))
        story.append(HRFlowable(width="100%", thickness=0.35, color=colors.HexColor("#CBD8E6"), spaceBefore=2, spaceAfter=2))

    story.append(Paragraph("DISPONIBILIDAD Y CONTACTO", heading))
    story.append(Paragraph(f'Disponible para desarrollo web, tecnología educativa, docencia de matemáticas y colaboración remota desde Lima, Perú.<br/>{pdf_link("juan@cabellosalirrosas.com", "mailto:juan@cabellosalirrosas.com")} · +51 925 475 034 · {pdf_link("TikTok educativo", "https://www.tiktok.com/@zenitmath")}', body))

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor(f"#{MUTED}"))
        canvas.drawCentredString(A4[0] / 2, 7 * mm, f"Juan Xavier Cabello Salirrosas · CV actualizado en octubre de 2026 · Página {document.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)


if __name__ == "__main__":
    build_docx()
    build_pdf()
    print(DOCX_OUT)
    print(PDF_OUT)
